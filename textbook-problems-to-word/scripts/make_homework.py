#!/usr/bin/env python3
"""Render manually verified PDF problem crops into a print-ready Word file."""

import json
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import pdfplumber
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def main(recipe_path):
    recipe = json.loads(Path(recipe_path).read_text(encoding="utf-8"))
    source = Path(recipe["pdf"])
    output = Path(recipe["output"])
    problems = recipe["problems"]
    if not problems:
        raise ValueError("Recipe contains no problems")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin, sec.bottom_margin = Inches(0.76), Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    title_style = doc.styles["Title"]
    title_style.font.name, title_style.font.size = "Arial", Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(11)
    title_properties = title_style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    doc.add_paragraph(recipe["title"], "Title")

    with pdfplumber.open(source) as pdf, TemporaryDirectory() as tmp:
        for index, problem in enumerate(problems):
            pieces = []
            for part_index, part in enumerate(problem.get("parts", [problem])):
                page_no = int(part["pdf_page"])
                if not 1 <= page_no <= len(pdf.pages):
                    raise ValueError(f"Invalid PDF page for {problem['number']}: {page_no}")
                page = pdf.pages[page_no - 1]
                left, top, right, bottom = map(float, part["crop"])
                if not (0 <= left < right <= page.width and 0 <= top < bottom <= page.height):
                    raise ValueError(f"Invalid crop for {problem['number']}")

                prefix = Path(tmp) / f"problem-{index+1}-part-{part_index+1}"
                subprocess.run([
                    "pdftoppm", "-f", str(page_no), "-l", str(page_no),
                    "-r", "450", "-png", "-singlefile", str(source), str(prefix)
                ], check=True, stdout=subprocess.DEVNULL)
                with Image.open(prefix.with_suffix(".png")) as image:
                    sx, sy = image.width / page.width, image.height / page.height
                    box = (round(left*sx), round(top*sy), round(right*sx), round(bottom*sy))
                    pieces.append(image.crop(box))

            gap = 40 if len(pieces) > 1 else 0
            crop = Image.new("RGB", (max(p.width for p in pieces),
                                     sum(p.height for p in pieces) + gap * (len(pieces)-1)), "white")
            offset = 0
            for piece in pieces:
                crop.paste(piece, ((crop.width-piece.width)//2, offset))
                offset += piece.height + gap
            crop_path = Path(tmp) / f"crop-{index+1}.png"
            crop.save(crop_path)

            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            if index:
                para.paragraph_format.page_break_before = True
            max_width, max_height = 6.25, 5.15
            width = min(max_width, max_height * crop.width / crop.height)
            para.add_run().add_picture(str(crop_path), width=Inches(width))

            if note := problem.get("note"):
                note_para = doc.add_paragraph()
                note_para.paragraph_format.space_before = Pt(9)
                label = note_para.add_run("Assignment note: ")
                label.bold = True
                note_para.add_run(note)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: make_homework.py recipe.json")
    main(sys.argv[1])
